import io
import math
from datetime import datetime

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    HRFlowable, PageBreak, KeepTogether,
)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

from . import models

import os as _os

# macOS ships Arial; Linux containers use Liberation Sans (metric-compatible, Cyrillic support).
if _os.path.isdir("/System/Library/Fonts/Supplemental"):
    _FONT_DIR   = "/System/Library/Fonts/Supplemental"
    _FONT_FILES = ("Arial.ttf", "Arial Bold.ttf", "Arial Italic.ttf", "Arial Bold Italic.ttf")
else:
    _FONT_DIR   = "/usr/share/fonts/truetype/liberation"
    _FONT_FILES = (
        "LiberationSans-Regular.ttf", "LiberationSans-Bold.ttf",
        "LiberationSans-Italic.ttf",  "LiberationSans-BoldItalic.ttf",
    )

pdfmetrics.registerFont(TTFont("Arial",            _os.path.join(_FONT_DIR, _FONT_FILES[0])))
pdfmetrics.registerFont(TTFont("Arial-Bold",       _os.path.join(_FONT_DIR, _FONT_FILES[1])))
pdfmetrics.registerFont(TTFont("Arial-Italic",     _os.path.join(_FONT_DIR, _FONT_FILES[2])))
pdfmetrics.registerFont(TTFont("Arial-BoldItalic", _os.path.join(_FONT_DIR, _FONT_FILES[3])))
pdfmetrics.registerFontFamily(
    "Arial",
    normal="Arial",
    bold="Arial-Bold",
    italic="Arial-Italic",
    boldItalic="Arial-BoldItalic",
)

_DIFF_ORDER = {models.Difficulty.B: 0, models.Difficulty.M: 1, models.Difficulty.H: 2}

_STRINGS = {
    "uk": {
        "module":        "Модуль",
        "created":       "Створено",
        "author":        "Автор",
        "section_b":     "БАЗОВИЙ РІВЕНЬ (Б)",
        "section_m":     "СЕРЕДНІЙ РІВЕНЬ (С)",
        "section_h":     "ВИЩИЙ РІВЕНЬ (В)",
        "diff_b":        "Б",
        "diff_m":        "С",
        "diff_h":        "В",
        "rubric_title":  "Рубрика оцінювання",
        "pass_line":     "Поріг зарахування: ≥ {thr} правильних з {n_bs} питань (Б+С)",
        "grade_col":     "Оцінка",
        "cond_col":      "Умова",
        "g5":            "5 — Відмінно",
        "g4":            "4 — Добре",
        "g3":            "3 — Задовільно",
        "g2":            "2 — Незадовільно",
        "g_pass":        "5 / 4 / 3 — Зараховано",
        "cond_pass":     "{bs_cond} правильно",
        "cond_fail":     "{bs_fail} правильно",
        "cond_all_v":    "{bs_cond}  та  {n_v}/{n_v} В правильно (100%)",
        "cond_some_v":   "{bs_cond}  та  1–{n_v1}/{n_v} В правильно (1–99%)",
        "cond_zero_v":   "{bs_cond}  та  0/{n_v} В правильно (0%)",
        "cond_1v_yes":   "{bs_cond}  та  1/1 В правильно",
        "cond_1v_no":    "{bs_cond}  та  0/1 В правильно",
        "cond_below":    "{bs_fail}",
        "note_no_v":     "У цьому тесті немає питань рівня В. Оцінка визначається лише за порогом {thr}/{n_bs} (Б+С).",
        "note_1v":       "Питання рівня В (1 шт.) не впливає на поріг зарахування. Правильна відповідь → «відмінно», неправильна → «задовільно».",
        "note_nv":       "Питання рівня В ({n_v} шт.) не впливають на поріг зарахування. Оцінка визначається відсотком правильних відповідей на питання В: 0% → «задовільно», 1–99% → «добре», 100% → «відмінно».",
        "note_rubric_docx": "Питання рівня В ({n_v} шт.) не впливають на поріг зарахування. Оцінка визначається відсотком правильних відповідей на питання В: 0% → «задовільно», 1–99% → «добре», 100% → «відмінно».",
        "rubric_heading":   "Рубрика оцінювання",
        "criteria_label":   "Критерії оцінювання:",
        "count_line":       "Базовий (Б): {n_b} пит.  |  Середній (С): {n_s} пит.  |  Вищий (В): {n_v} пит.  |  Разом: {total} питань",
        "pass_line_docx":   "Поріг зарахування: ≥ {thr} правильних з {n_bs} (Б+С)",
        "pass_threshold_policy": "параметр навчального закладу",
        "note_prefix":      "Примітка: ",
    },
    "en": {
        "module":        "Module",
        "created":       "Created",
        "author":        "Author",
        "section_b":     "BASIC LEVEL (B)",
        "section_m":     "INTERMEDIATE LEVEL (M)",
        "section_h":     "ADVANCED LEVEL (H)",
        "diff_b":        "B",
        "diff_m":        "M",
        "diff_h":        "H",
        "rubric_title":  "Scoring Rubric",
        "pass_line":     "Pass threshold: ≥ {thr} correct out of {n_bs} questions (B+M)",
        "grade_col":     "Grade",
        "cond_col":      "Condition",
        "g5":            "5 — Excellent",
        "g4":            "4 — Good",
        "g3":            "3 — Satisfactory",
        "g2":            "2 — Fail",
        "g_pass":        "5 / 4 / 3 — Pass",
        "cond_pass":     "{bs_cond} correct",
        "cond_fail":     "{bs_fail} correct",
        "cond_all_v":    "{bs_cond}  and  {n_v}/{n_v} H correct (100%)",
        "cond_some_v":   "{bs_cond}  and  1–{n_v1}/{n_v} H correct (1–99%)",
        "cond_zero_v":   "{bs_cond}  and  0/{n_v} H correct (0%)",
        "cond_1v_yes":   "{bs_cond}  and  1/1 H correct",
        "cond_1v_no":    "{bs_cond}  and  0/1 H correct",
        "cond_below":    "{bs_fail}",
        "note_no_v":     "This test has no advanced-level (H) questions. The grade is determined solely by the {thr}/{n_bs} (B+M) threshold.",
        "note_1v":       "The single advanced-level (H) question does not affect the pass threshold. Correct → 'Excellent', incorrect → 'Satisfactory'.",
        "note_nv":       "Advanced-level (H) questions ({n_v}) do not affect the pass threshold. The grade is determined by the percentage of H questions answered correctly: 0% → 'Satisfactory', 1–99% → 'Good', 100% → 'Excellent'.",
        "note_rubric_docx": "Advanced-level (H) questions ({n_v}) do not affect the pass threshold. Grade is based on % of H questions correct: 0% → satisfactory, 1–99% → good, 100% → excellent.",
        "rubric_heading":   "Scoring Rubric",
        "criteria_label":   "Grading criteria:",
        "count_line":       "Basic (B): {n_b} q.  |  Intermediate (M): {n_s} q.  |  Advanced (H): {n_v} q.  |  Total: {total} questions",
        "pass_line_docx":   "Pass threshold: ≥ {thr} correct out of {n_bs} (B+M)",
        "pass_threshold_policy": "institutional policy",
        "note_prefix":      "Note: ",
    },
}


def _org_header(org: dict) -> str:
    name = (org or {}).get("org_name", "").strip()
    dept = (org or {}).get("org_dept", "").strip()
    if name and dept:
        return f"{name} · {dept}"
    return name or dept or "ОНМедУ · Мікробіологія"


def export_test_docx(test: models.Test, org: dict = None, lang: str = "uk", pass_pct: int = 75) -> bytes:
    s = _STRINGS.get(lang, _STRINGS["uk"])
    doc = Document()

    title = doc.add_heading(test.name, level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    module     = test.module
    discipline = module.discipline
    p = doc.add_paragraph()
    p.add_run(f"{discipline.name}").bold = True
    doc.add_paragraph(f"{s['module']} {module.number}: {module.name}")
    doc.add_paragraph()

    def _eff_diff(q):
        return q.final_difficulty or q.proposed_difficulty

    questions = sorted(
        [tq.question for tq in test.test_questions],
        key=lambda q: _DIFF_ORDER.get(_eff_diff(q), 99),
    )

    counts = {models.Difficulty.B: 0, models.Difficulty.M: 0, models.Difficulty.H: 0}
    for q in questions:
        d = _eff_diff(q)
        if d in counts:
            counts[d] += 1

    diff_section = {
        models.Difficulty.B: s["section_b"].title(),
        models.Difficulty.M: s["section_m"].title(),
        models.Difficulty.H: s["section_h"].title(),
    }

    current_section = None
    for q in questions:
        section = diff_section.get(_eff_diff(q))
        if section != current_section:
            current_section = section
            h = doc.add_paragraph()
            r = h.add_run(section)
            r.bold = True
            r.font.size = Pt(11)

        p = doc.add_paragraph(style="List Number")
        p.add_run(q.text)

        if q.model_answer:
            ans = doc.add_paragraph()
            run = ans.add_run(f"  ✓ {q.model_answer}")
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x33, 0x77, 0x33)
            run.italic = True

    # ── Scoring rubric ──────────────────────────────────────
    doc.add_page_break()
    doc.add_heading(s["rubric_heading"], level=2)

    n_b  = counts[models.Difficulty.B]
    n_s  = counts[models.Difficulty.M]
    n_v  = counts[models.Difficulty.H]
    n_bs = n_b + n_s
    thr  = math.ceil(n_bs * pass_pct / 100)

    doc.add_paragraph(s["count_line"].format(n_b=n_b, n_s=n_s, n_v=n_v, total=n_b+n_s+n_v))
    policy_label = s["pass_threshold_policy"]
    method_note  = (org or {}).get("pass_threshold_method", "").strip()
    method_str   = f" ({method_note})" if method_note else ""
    doc.add_paragraph(s["pass_line_docx"].format(thr=thr, n_bs=n_bs) + f"  [{policy_label}{method_str}]")
    doc.add_paragraph()

    grading = doc.add_paragraph()
    grading.add_run(s["criteria_label"]).bold = True

    bs_cond = f"≥ {thr}/{n_bs}"
    bs_fail = f"< {thr}/{n_bs}"

    if n_v == 0:
        rows = [
            (s["g_pass"],  s["cond_pass"].format(bs_cond=bs_cond)),
            (s["g2"],      s["cond_fail"].format(bs_fail=bs_fail)),
        ]
    elif n_v == 1:
        rows = [
            (s["g5"], s["cond_1v_yes"].format(bs_cond=bs_cond)),
            (s["g3"], s["cond_1v_no"].format(bs_cond=bs_cond)),
            (s["g2"], s["cond_below"].format(bs_fail=bs_fail)),
        ]
    else:
        rows = [
            (s["g5"], s["cond_all_v"].format(bs_cond=bs_cond, n_v=n_v)),
            (s["g4"], s["cond_some_v"].format(bs_cond=bs_cond, n_v=n_v, n_v1=n_v-1)),
            (s["g3"], s["cond_zero_v"].format(bs_cond=bs_cond, n_v=n_v)),
            (s["g2"], s["cond_below"].format(bs_fail=bs_fail)),
        ]

    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].paragraphs[0].add_run(s["grade_col"]).bold = True
    hdr[1].paragraphs[0].add_run(s["cond_col"]).bold = True
    for label, condition in rows:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = condition

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.add_run(s["note_prefix"]).bold = True
    if n_v == 0:
        note.add_run(s["note_no_v"].format(thr=thr, n_bs=n_bs))
    elif n_v == 1:
        note.add_run(s["note_1v"])
    else:
        note.add_run(s["note_rubric_docx"].format(n_v=n_v))

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


# ── Design tokens mirroring app.css ────────────────────────────────────────
_INK        = colors.HexColor("#0f172a")   # --text
_INK_2      = colors.HexColor("#334155")   # --text-2
_INK_4      = colors.HexColor("#94a3b8")   # --text-4
_ACCENT     = colors.HexColor("#0ea5e9")   # --accent
_SURFACE    = colors.HexColor("#f8fafc")   # --surface
_BORDER     = colors.HexColor("#e2e8f0")   # --border
_DIFF_B     = colors.HexColor("#2563eb")   # diff-b
_DIFF_S     = colors.HexColor("#d97706")   # diff-s
_DIFF_V     = colors.HexColor("#dc2626")   # diff-v
_GREEN      = colors.HexColor("#16a34a")

_W, _H = A4
_ML, _MR, _MT, _MB = 20*mm, 20*mm, 22*mm, 22*mm


def _diff_color(d):
    return {models.Difficulty.B: _DIFF_B, models.Difficulty.M: _DIFF_S, models.Difficulty.H: _DIFF_V}.get(d, _INK)


def _diff_label(d, s):
    return {
        models.Difficulty.B: s["diff_b"],
        models.Difficulty.M: s["diff_m"],
        models.Difficulty.H: s["diff_h"],
    }.get(d, "?")


def _build_styles():
    base = getSampleStyleSheet()
    def s(name, **kw):
        return ParagraphStyle(name, **kw)

    return {
        "title": s("title",
            fontName="Arial-Bold", fontSize=17, leading=22,
            textColor=_INK, alignment=TA_LEFT, spaceAfter=2),
        "subtitle": s("subtitle",
            fontName="Arial", fontSize=9, leading=13,
            textColor=_INK_4, spaceAfter=10),
        "section": s("section",
            fontName="Arial-Bold", fontSize=7.5, leading=11,
            textColor=_INK_4, spaceBefore=10, spaceAfter=4,
            letterSpacing=0.9),
        "q_num": s("q_num",
            fontName="Arial-Bold", fontSize=8.5, leading=12,
            textColor=_INK_4),
        "q_text": s("q_text",
            fontName="Arial", fontSize=9.5, leading=14,
            textColor=_INK),
        "q_answer": s("q_answer",
            fontName="Arial-Italic", fontSize=8.5, leading=12,
            textColor=_GREEN),
        "rubric_title": s("rubric_title",
            fontName="Arial-Bold", fontSize=11, leading=15,
            textColor=_INK, spaceBefore=6, spaceAfter=6),
        "rubric_body": s("rubric_body",
            fontName="Arial", fontSize=8.5, leading=13,
            textColor=_INK_2, spaceAfter=3),
        "note": s("note",
            fontName="Arial", fontSize=8, leading=12,
            textColor=_INK_4, spaceBefore=8),
        "footer": s("footer",
            fontName="Arial", fontSize=7.5, leading=10,
            textColor=_INK_4, alignment=TA_CENTER),
    }


class _HeaderFooterCanvas(canvas.Canvas):
    """Draws header rule + footer on every page."""

    def __init__(self, *args, test_name="", created_at="", org_header="", **kwargs):
        super().__init__(*args, **kwargs)
        self._test_name = test_name
        self._created_at = created_at
        self._org_header = org_header
        self._saved_state = []

    def showPage(self):
        self._saved_state.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        page_count = len(self._saved_state)
        for i, state in enumerate(self._saved_state, 1):
            self.__dict__.update(state)
            self._draw_chrome(i, page_count)
            super().showPage()
        super().save()

    def _draw_chrome(self, page_num, total):
        self.saveState()
        w = _W - _ML - _MR

        # top rule
        self.setStrokeColor(_ACCENT)
        self.setLineWidth(2)
        self.line(_ML, _H - _MT + 4*mm, _ML + w, _H - _MT + 4*mm)

        self.setFont("Arial", 7.5)
        self.setFillColor(_INK_4)
        self.drawRightString(_ML + w, _H - _MT + 5.5*mm, self._org_header)

        # bottom rule
        self.setStrokeColor(_BORDER)
        self.setLineWidth(0.5)
        self.line(_ML, _MB - 4*mm, _ML + w, _MB - 4*mm)

        # footer left: test name
        self.setFont("Arial", 7.5)
        self.setFillColor(_INK_4)
        self.drawString(_ML, _MB - 6.5*mm, self._test_name)

        # footer right: page n/total + date
        right_text = f"{self._created_at}   {page_num} / {total}"
        self.drawRightString(_ML + w, _MB - 6.5*mm, right_text)

        self.restoreState()


def export_test_pdf(test: models.Test, org: dict = None, lang: str = "uk", pass_pct: int = 75) -> bytes:
    s = _STRINGS.get(lang, _STRINGS["uk"])
    buf = io.BytesIO()
    module     = test.module
    discipline = module.discipline
    created_str = test.created_at.strftime("%d.%m.%Y")
    org_str = _org_header(org)

    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=_ML, rightMargin=_MR,
        topMargin=_MT, bottomMargin=_MB,
        title=test.name,
        author=org_str,
    )

    ST = _build_styles()
    story = []

    # ── Cover header ────────────────────────────────────────
    story.append(Spacer(1, 2*mm))
    story.append(Paragraph(test.name, ST["title"]))
    meta_parts = [
        discipline.name,
        f"{s['module']} {module.number}: {module.name}",
        f"{s['created']}: {test.created_at.strftime('%d.%m.%Y %H:%M')}",
        f"{s['author']}: {test.creator.name}",
    ]
    story.append(Paragraph("  ·  ".join(meta_parts), ST["subtitle"]))
    story.append(HRFlowable(width="100%", thickness=0.5, color=_BORDER, spaceAfter=8))

    # ── Sort questions ───────────────────────────────────────
    def _eff_diff(q):
        return q.final_difficulty or q.proposed_difficulty

    questions = sorted(
        [tq.question for tq in test.test_questions],
        key=lambda q: _DIFF_ORDER.get(_eff_diff(q), 99),
    )
    counts = {models.Difficulty.B: 0, models.Difficulty.M: 0, models.Difficulty.H: 0}
    for q in questions:
        d = _eff_diff(q)
        if d in counts:
            counts[d] += 1

    # ── Questions ────────────────────────────────────────────
    section_labels = {
        models.Difficulty.B: s["section_b"],
        models.Difficulty.M: s["section_m"],
        models.Difficulty.H: s["section_h"],
    }
    current_section = None
    q_number = 1

    for q in questions:
        sec = section_labels.get(_eff_diff(q))
        if sec != current_section:
            current_section = sec
            story.append(Spacer(1, 3*mm))
            pill_color = _diff_color(_eff_diff(q))
            section_table = Table(
                [[Paragraph(sec, ParagraphStyle("sh",
                    fontName="Arial-Bold", fontSize=7, leading=9,
                    textColor=colors.white))]],
                colWidths=[55*mm],
            )
            section_table.setStyle(TableStyle([
                ("BACKGROUND", (0,0), (-1,-1), pill_color),
                ("ROUNDEDCORNERS", [3]),
                ("TOPPADDING",  (0,0), (-1,-1), 3),
                ("BOTTOMPADDING", (0,0), (-1,-1), 3),
                ("LEFTPADDING",  (0,0), (-1,-1), 6),
                ("RIGHTPADDING", (0,0), (-1,-1), 6),
            ]))
            story.append(section_table)
            story.append(Spacer(1, 3*mm))

        dc = _diff_color(_eff_diff(q))
        num_para  = Paragraph(f"<b>{q_number}.</b>", ST["q_num"])
        text_para = Paragraph(q.text, ST["q_text"])

        cell_content = [text_para]
        if q.model_answer:
            cell_content.append(Spacer(1, 1*mm))
            cell_content.append(Paragraph(f"✓ {q.model_answer}", ST["q_answer"]))

        badge = Table(
            [[Paragraph(f"<b>{_diff_label(_eff_diff(q), s)}</b>",
                ParagraphStyle("badge",
                    fontName="Arial-Bold", fontSize=7.5, leading=9,
                    textColor=dc, alignment=TA_CENTER))]],
            colWidths=[5.5*mm],
        )
        badge.setStyle(TableStyle([
            ("BOX", (0,0), (-1,-1), 0.5, dc),
            ("TOPPADDING",  (0,0), (-1,-1), 2),
            ("BOTTOMPADDING", (0,0), (-1,-1), 2),
            ("LEFTPADDING",  (0,0), (-1,-1), 1),
            ("RIGHTPADDING", (0,0), (-1,-1), 1),
        ]))

        row = Table(
            [[num_para, cell_content, badge]],
            colWidths=[8*mm, None, 8*mm],
        )
        row.setStyle(TableStyle([
            ("VALIGN",       (0,0), (-1,-1), "TOP"),
            ("LEFTPADDING",  (0,0), (-1,-1), 0),
            ("RIGHTPADDING", (0,0), (-1,-1), 0),
            ("TOPPADDING",   (0,0), (-1,-1), 5),
            ("BOTTOMPADDING",(0,0), (-1,-1), 5),
            ("LINEBELOW",    (0,0), (-1,-1), 0.3, _BORDER),
        ]))

        story.append(KeepTogether(row))
        q_number += 1

    # ── Rubric ───────────────────────────────────────────────
    story.append(Spacer(1, 8*mm))
    story.append(Paragraph(s["rubric_title"], ST["rubric_title"]))
    story.append(HRFlowable(width="100%", thickness=0.5, color=_BORDER, spaceAfter=6))

    n_b  = counts[models.Difficulty.B]
    n_s  = counts[models.Difficulty.M]
    n_v  = counts[models.Difficulty.H]
    n_bs = n_b + n_s
    thr  = math.ceil(n_bs * pass_pct / 100)
    policy_label = s["pass_threshold_policy"]
    method_note  = (org or {}).get("pass_threshold_method", "").strip()
    method_str   = f" ({method_note})" if method_note else ""

    story.append(Paragraph(
        s["pass_line"].format(thr=f"<b>{thr}</b>", n_bs=n_bs) + f"  <font color='#94a3b8' size='7'>[{policy_label}{method_str}]</font>",
        ST["rubric_body"],
    ))
    story.append(Spacer(1, 5*mm))

    bs_cond = f"≥ {thr}/{n_bs}"
    bs_fail = f"< {thr}/{n_bs}"

    if n_v == 0:
        grade_rows = [
            (s["g_pass"], s["cond_pass"].format(bs_cond=bs_cond)),
            (s["g2"],     s["cond_fail"].format(bs_fail=bs_fail)),
        ]
        grade_colors = [
            ("BACKGROUND", (0,1), (0,1), colors.HexColor("#dcfce7")),
            ("BACKGROUND", (0,2), (0,2), colors.HexColor("#fee2e2")),
        ]
        note_text = s["note_no_v"].format(thr=thr, n_bs=n_bs)
    elif n_v == 1:
        grade_rows = [
            (s["g5"], s["cond_1v_yes"].format(bs_cond=bs_cond)),
            (s["g3"], s["cond_1v_no"].format(bs_cond=bs_cond)),
            (s["g2"], s["cond_below"].format(bs_fail=bs_fail)),
        ]
        grade_colors = [
            ("BACKGROUND", (0,1), (0,1), colors.HexColor("#dcfce7")),
            ("BACKGROUND", (0,2), (0,2), colors.HexColor("#fef9c3")),
            ("BACKGROUND", (0,3), (0,3), colors.HexColor("#fee2e2")),
        ]
        note_text = s["note_1v"]
    else:
        grade_rows = [
            (s["g5"], s["cond_all_v"].format(bs_cond=bs_cond, n_v=n_v)),
            (s["g4"], s["cond_some_v"].format(bs_cond=bs_cond, n_v=n_v, n_v1=n_v-1)),
            (s["g3"], s["cond_zero_v"].format(bs_cond=bs_cond, n_v=n_v)),
            (s["g2"], s["cond_below"].format(bs_fail=bs_fail)),
        ]
        grade_colors = [
            ("BACKGROUND", (0,1), (0,1), colors.HexColor("#dcfce7")),
            ("BACKGROUND", (0,2), (0,2), colors.HexColor("#fef9c3")),
            ("BACKGROUND", (0,3), (0,3), colors.HexColor("#fef9c3")),
            ("BACKGROUND", (0,4), (0,4), colors.HexColor("#fee2e2")),
        ]
        note_text = s["note_nv"].format(n_v=n_v)

    grades = [
        [Paragraph(f"<b>{s['grade_col']}</b>", ST["rubric_body"]),
         Paragraph(f"<b>{s['cond_col']}</b>",  ST["rubric_body"])],
    ] + [
        [Paragraph(label, ST["rubric_body"]), Paragraph(cond, ST["rubric_body"])]
        for label, cond in grade_rows
    ]

    grade_style = [
        ("BACKGROUND",    (0,0), (-1,0), _SURFACE),
        ("LINEBELOW",     (0,0), (-1,0), 0.5, _BORDER),
        ("BOX",           (0,0), (-1,-1), 0.5, _BORDER),
        ("INNERGRID",     (0,0), (-1,-1), 0.3, _BORDER),
        ("TOPPADDING",    (0,0), (-1,-1), 5),
        ("BOTTOMPADDING", (0,0), (-1,-1), 5),
        ("LEFTPADDING",   (0,0), (-1,-1), 8),
    ] + grade_colors

    grade_table = Table(grades, colWidths=[52*mm, None])
    grade_table.setStyle(TableStyle(grade_style))
    story.append(grade_table)
    story.append(Paragraph(note_text, ST["note"]))

    # ── Build ────────────────────────────────────────────────
    def make_canvas(filename, **kwargs):
        return _HeaderFooterCanvas(
            filename,
            test_name=test.name,
            created_at=created_str,
            org_header=org_str,
            **kwargs,
        )

    doc.build(story, canvasmaker=make_canvas)
    buf.seek(0)
    return buf.read()
